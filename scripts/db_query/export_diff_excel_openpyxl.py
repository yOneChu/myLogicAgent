import json
import os
import csv
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

os.makedirs("output_excel", exist_ok=True)
os.makedirs("output_csv", exist_ok=True)

with open("output_csv/EL_PA103A03_v44_v45_korean_diff.json", "r", encoding="utf-8") as f:
    diff_data = json.load(f)

excel_rows = []

for item in diff_data:
    no = item["NO"]
    t = item["TYPE"]
    
    if t == "MODIFIED":
        changed_fields = list(item.get("field_diffs", {}).keys())
        detail_list = []
        for fld, vals in item.get("field_diffs", {}).items():
            detail_list.append(f"{fld}: [{vals['v44']}] -> [{vals['v45']}]")
            
        excel_rows.append({
            "NO": no,
            "구분": "수정(MODIFIED)",
            "주요 변경 필드": ", ".join(changed_fields),
            "v44 정보": item.get("v44_summary"),
            "v45 정보": item.get("v45_summary"),
            "상세 변경 내용": " | ".join(detail_list)
        })
    elif t == "ADDED_IN_V45":
        excel_rows.append({
            "NO": no,
            "구분": "v45 신규 추가",
            "주요 변경 필드": "전체 행 추가",
            "v44 정보": "-",
            "v45 정보": item.get("v45_summary"),
            "상세 변경 내용": "v45 버전에서 신규 추가된 행"
        })
    elif t == "DELETED_IN_V45":
        excel_rows.append({
            "NO": no,
            "구분": "v45 삭제",
            "주요 변경 필드": "전체 행 삭제",
            "v44 정보": item.get("v44_summary"),
            "v45 정보": "-",
            "상세 변경 내용": "v45 버전에서 삭제된 행"
        })

wb = Workbook()
ws = wb.active
ws.title = "EL_PA103A03_v44_v45_Diff"

headers = ["NO", "구분", "주요 변경 필드", "v44 정보", "v45 정보", "상세 변경 내용"]
ws.append(headers)

# Styling
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

row_font = Font(name="맑은 고딕", size=10)
for row_idx, data in enumerate(excel_rows, start=2):
    row_vals = [data[h] for h in headers]
    ws.append(row_vals)
    for col_idx, cell in enumerate(ws[row_idx], start=1):
        cell.font = row_font
        cell.border = thin_border
        if col_idx in (1, 2):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 50
ws.column_dimensions['E'].width = 50
ws.column_dimensions['F'].width = 60

excel_path = "output_excel/EL_PA103A03_v44_v45_diff_report.xlsx"
wb.save(excel_path)
print(f"Excel report successfully generated at: {excel_path}")
