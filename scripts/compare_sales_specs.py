import json
import ssl
import sys
import os
import csv
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

API_BASE_URL = "https://vault-in.hdel.co.kr:8070/api"
API_KEY = "subae"

def fetch_json(endpoint: str, params: dict = None, insecure: bool = True) -> any:
    if params is None:
        params = {}
    params["key"] = API_KEY
    url = f"{API_BASE_URL}/{endpoint}?{urlencode(params)}"
    
    req = Request(
        url,
        method="GET",
        headers={
            "Accept": "application/json",
            "User-Agent": "sales-spec-compare/2.0",
        },
    )
    
    context = ssl._create_unverified_context() if insecure else None
    
    with urlopen(req, timeout=120, context=context) as response:
        raw_body = response.read()
        
    try:
        body = raw_body.decode("utf-8")
    except UnicodeDecodeError:
        body = raw_body.decode("cp949", errors="replace")
        
    return json.loads(body)

def get_code_field_map():
    try:
        res = fetch_json("getCodeField")
        field_map = {}
        items = res.get("data", res) if isinstance(res, dict) else res
        if isinstance(items, list):
            for item in items:
                name = item.get("NAME") or item.get("code")
                tit = item.get("TIT") or item.get("codeName")
                if name and tit:
                    field_map[name.upper()] = tit
        return field_map
    except Exception as e:
        print(f"Warning: Failed to fetch code fields: {e}", file=sys.stderr)
        return {}

def main():
    field_map = get_code_field_map()
    print(f"Loaded {len(field_map)} field title mappings.")

    sql_raw = """
    SELECT V.*
    FROM HDEL_DEFAULT.ELV_INFO$VF V,
         HDEL_DEFAULT.ELV_INFO$ID A
    WHERE V.vf$identity = A.id$ouid
      AND V.vf$ouid = A.id$wip
      AND V.MD$NUMBER IN ('213249L36', '210003L45')
    """
    
    sql_decoded = """
    SELECT 
        V.MD$NUMBER AS PRODUCTNO,
        HDEL_DEFAULT.COD(V.EL_AOPEN) AS EL_AOPEN_DEC,
        HDEL_DEFAULT.CODN(V.EL_AUSE) AS EL_AUSE_DEC,
        HDEL_DEFAULT.CODN(V.EL_ABRAND) AS EL_ABRAND_DEC,
        HDEL_DEFAULT.CODN(V.EL_ATYP) AS EL_ATYP_DEC,
        HDEL_DEFAULT.CODN(V.EL_ASPD) AS EL_ASPD_DEC,
        HDEL_DEFAULT.CODN(V.EL_ACAPA) AS EL_ACAPA_DEC,
        HDEL_DEFAULT.COD(V.EL_ECWRL) AS EL_ECWRL_DEC,
        HDEL_DEFAULT.COD(V.EL_ETM) AS EL_ETM_DEC,
        HDEL_DEFAULT.COD(V.EL_ECSF) AS EL_ECSF_DEC,
        HDEL_DEFAULT.COD(V.EL_ASPC) AS EL_ASPC_DEC,
        HDEL_DEFAULT.COD(V.EL_ASPCD) AS EL_ASPCD_DEC,
        HDEL_DEFAULT.COD(V.EL_BCL) AS EL_BCL_DEC,
        HDEL_DEFAULT.COD(V.EL_DCRG) AS EL_DCRG_DEC,
        HDEL_DEFAULT.COD(V.EL_ASPSCD) AS EL_ASPSCD_DEC,
        HDEL_DEFAULT.COD(V.EL_ASPSC) AS EL_ASPSC_DEC,
        HDEL_DEFAULT.COD(V.EL_BCDM) AS EL_BCDM_DEC,
        HDEL_DEFAULT.COD(V.EL_BWALLT) AS EL_BWALLT_DEC,
        HDEL_DEFAULT.COD(V.EL_BCLCDL) AS EL_BCLCDL_DEC,
        HDEL_DEFAULT.COD(V.EL_BMOPB) AS EL_BMOPB_DEC,
        HDEL_DEFAULT.COD(V.EL_BETM) AS EL_BETM_DEC,
        HDEL_DEFAULT.COD(V.EL_BOPBSWD) AS EL_BOPBSWD_DEC
    FROM HDEL_DEFAULT.ELV_INFO$VF V,
         HDEL_DEFAULT.ELV_INFO$ID A
    WHERE V.vf$identity = A.id$ouid
      AND V.vf$ouid = A.id$wip
      AND V.MD$NUMBER IN ('213249L36', '210003L45')
    """

    print("Fetching raw data...")
    raw_res = fetch_json("executeQuery", {"sql": sql_raw})
    
    print("Fetching decoded spec values...")
    dec_res = fetch_json("executeQuery", {"sql": sql_decoded})

    hogi1 = '213249L36'
    hogi2 = '210003L45'

    dict_raw = {item.get('MD$NUMBER') or item.get('PRODUCTNO'): item for item in raw_res}
    dict_dec = {item.get('PRODUCTNO') or item.get('MD$NUMBER'): item for item in dec_res}

    row1_raw = dict_raw.get(hogi1, {})
    row2_raw = dict_raw.get(hogi2, {})
    
    row1_dec = dict_dec.get(hogi1, {})
    row2_dec = dict_dec.get(hogi2, {})

    all_keys = set(list(row1_raw.keys()) + list(row2_raw.keys()))
    
    meta_title_map = {
        'MD$DESC': '수주명 (현장명)',
        'MD$NUMBER': '호기번호',
        'MD$USER': '등록자 사번',
        'MD$CDATE': '생성/등록 일시',
        'MD$MDATE': '수정 일시',
        'MD$STATUS': '진행 상태',
        'MANAGER_E': '전기 담당자',
        'MANAGER_M': '기계 담당자',
        'MANAGER2': '영업/기타 담당자',
        'SAVEUSER': '최종 저장자',
        'PHONE': '연락처1',
        'PHONE2': '연락처2',
        'PARTNOGROUP': '자재그룹번호',
        'VF$IDENTITY': 'Identity ID',
        'VF$OUID': 'WIP OUID',
        'VF$VERSION': '버전',
        'ERPTPDATE': 'ERP 전송일시',
    }

    diffs = []
    ignore_keys = {'VF$IDENTITY', 'VF$OUID', 'VF$WIP', 'ID$OUID', 'ID$WIP', 'MD$ID'}

    for key in sorted(all_keys):
        if key in ignore_keys:
            continue
            
        val1 = str(row1_raw.get(key, '')).strip() if row1_raw.get(key) is not None else ''
        val2 = str(row2_raw.get(key, '')).strip() if row2_raw.get(key) is not None else ''
        
        if val1 != val2:
            title = meta_title_map.get(key) or field_map.get(key.upper(), '')
            dec_key = f"{key}_DEC"
            dec1 = str(row1_dec.get(dec_key, '')).strip() if row1_dec.get(dec_key) is not None else ''
            dec2 = str(row2_dec.get(dec_key, '')).strip() if row2_dec.get(dec_key) is not None else ''
            
            diffs.append({
                'COLUMN': key,
                'TITLE': title,
                'VAL_213249L36': val1,
                'VAL_210003L45': val2,
                'DEC_213249L36': dec1,
                'DEC_210003L45': dec2
            })

    os.makedirs('output_csv', exist_ok=True)
    os.makedirs('output_excel', exist_ok=True)

    # 1. Export CSV
    csv_path = os.path.join('output_csv', 'sales_spec_diff_213249L36_210003L45.csv')
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['COLUMN', 'TITLE', 'VAL_213249L36', 'VAL_210003L45', 'DEC_213249L36', 'DEC_210003L45'])
        writer.writeheader()
        writer.writerows(diffs)

    # 2. Export Excel
    excel_path = os.path.join('output_excel', 'sales_spec_diff_213249L36_210003L45.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = "영업사양 비교"

    headers = ['사양 코드 (COLUMN)', '사양 명칭 (TITLE)', f'호기 213249L36', f'호기 210003L45', '213249L36 명칭(디코딩)', '210003L45 명칭(디코딩)']
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for diff in diffs:
        row = [
            diff['COLUMN'],
            diff['TITLE'],
            diff['VAL_213249L36'],
            diff['VAL_210003L45'],
            diff['DEC_213249L36'],
            diff['DEC_210003L45']
        ]
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=len(diffs)+1, min_col=1, max_col=6):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10)
            cell.border = thin_border

    wb.save(excel_path)

    json_path = os.path.join('output_csv', 'diff_result.json')
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(diffs, f, ensure_ascii=False, indent=2)

    print(f"Comparison complete. Found {len(diffs)} differences.")
    print(f"CSV exported to: {csv_path}")
    print(f"Excel exported to: {excel_path}")

if __name__ == '__main__':
    main()
